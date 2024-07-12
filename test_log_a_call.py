import time

from openpyxl import load_workbook
from selenium.webdriver import ActionChains

from Pages.Keen_Task import New_Task
from utilities import logger_utils
from Pages.Keen_Log_a_call import New_Log
from tests.test_ALogin import Test_login
import pytest
import allure

from utilities.XLUtilities import Excel_data
workbook = load_workbook('../test_data/Keen_Testdata.xlsx')
path="../test_data/Keen_Testdata.xlsx"
readdata_tasks=Excel_data(path, "Tasks")
sheet_log_a_call=workbook['Call_logs']
class Test_log_a_call():
    def __init__(self,driver):
        self.driver=driver
        self.log = logger_utils.get_logger(self.driver)
        self.action = ActionChains(self.driver)
    ##Test Case Description: Create and Assosciate A Call to a Particular Member
    ##TestCase Steps:1.Login to SFDC--->Click on Keen Leads and Members--->Click on log a call
    ##Create New Log and Assosicate to a Selected Member
    def test_log_a_call(self,row_number,start_date,end_date):
        Subject = readdata_tasks.readdata(row_number, 1)
        Status = readdata_tasks.readdata(row_number, 2)
        Dispositions = readdata_tasks.readdata(row_number, 3)
        Start_Time = readdata_tasks.readdata(row_number, 5)
        End_Time = readdata_tasks.readdata(row_number, 7)
        Score = readdata_tasks.readdata(row_number, 8)
        Task_Type = readdata_tasks.readdata(row_number, 9)
        Task_Subtype = readdata_tasks.readdata(row_number, 10)
        keen = New_Log(self.driver)
        self.driver.execute_script("scroll(0, 250);")
        keen.log_a_call()
        self.log.info("***** Fill the form as per the test data *****")
        keen.subject("Meeting")
        list_oif_available = keen.dispositions()
        for dispositions in list_oif_available:
            if dispositions.text == "Completed app":
                self.action.move_to_element(dispositions).click().perform()
                break
        keen.select_satisfaction_score()
        list_of_score = keen.select_status()
        for score in list_of_score:
            if score.text == Score:
                self.action.move_to_element(score).click().perform()
                break
        keen.start_date(start_date)
        keen.End_date(end_date)
        keen.status()
        list_of_status = keen.select_status()
        for status in list_of_status:
            if status.text == Status:
                self.action.move_to_element(status).click().perform()
                break
        keen.save()
        self.log.info("***** Log a call is saved ***** ")
    def test_validation_log_a_call(self):
        keen=New_Log(self.driver)
        self.log.info("***** Validation is started ***** ")
        try:
            validation_of_data = keen.validation_of_data()
            for i in validation_of_data:
                if i.text == "Other":
                    assert True
                elif i.text == "Open":
                    assert True
                elif i.text == "12/16/2022, 12:00 PM ":
                    assert True
                elif i.text == "12/17/2022, 12:00 PM":
                    assert True
                elif i.text == "Completed app":
                    assert True
                else:
                    assert False
        except Exception as ex:
            print(ex)
        self.log.info("***** Validation is completed ***** ")


    def test_log_call(self,row_number):
        keen = New_Log(self.driver)
        self.driver.execute_script("scroll(0, 250);")
        keen.log_a_call()
        print("Row number:", row_number)
        data_row = sheet_log_a_call[row_number]
        column_field_mapping = {
            'Subject': 'Subject',
            'Dispositions': 'Dispositions',
            'Satisfaction_Score': 'Satisfaction_Score',
            'Start_Date': 'Start_Date',
            'Start_time': 'Start_time', 'End_Date': 'End_Date', 'Status': 'Status',
            'End_Time': 'End_Time',
        }
        field_locators = {
            'Subject': (New_Log.log_subject),
            'Dispositions': (New_Log.call_dispositions),
            'Satisfaction_Score': (New_Log.log_satisfaction_score),
            'Start_Date': (New_Log.log_start_date),
            'Start_time': (New_Log.log_start_time),
            'End_Date': (New_Log.log_end_date),
            'Status': (New_Log.log_status),
            'End_Time': (New_Log.log_end_time),
        }
        for cell in data_row:
            column_name = sheet_log_a_call.cell(row=1, column=cell.column).value
            print("Column name:", column_name)
            if column_name in column_field_mapping:
                field_name = column_field_mapping[column_name]
                print("mapping is done")
                column_value = cell.value
                print("input:",column_value)
                locator = field_locators[field_name]
                element = self.driver.find_element(*locator)
                if element.tag_name == 'div':
                    if field_name == 'Status':
                        list_of_status = keen.select_status()
                        for status in list_of_status:
                            if status.text == column_value:
                                self.action.move_to_element(status).click().perform()
                                break
                    if field_name == 'Satisfaction_Score':
                        element.click()
                        list_score=keen.select_satisfaction_score()
                        for score in list_score:
                            if score.text == column_value:
                                self.driver.execute_script("arguments[0].click();", score)
                                break
                    if field_name == 'Dispositions':
                        list_oif_available = keen.dispositions()
                        for dispositions in list_oif_available:
                            if dispositions.text == column_value:
                                self.driver.execute_script("arguments[0].click();", dispositions)
                                keen.click_move_to_choosen()
                                break
                else:
                    element.clear()
                    element.send_keys(column_value)






