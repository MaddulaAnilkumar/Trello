import time
from datetime import datetime
import pytest
import pytz
import allure
from selenium.webdriver import ActionChains
from selenium.webdriver.common.by import By
from selenium.webdriver.support.wait import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from tests.test_create_Capture_full_Member_details import Test_create_capture_full_member_details
from utilities.XLUtilities import Excel_data
from utilities.action_utils import ActionUtils
from openpyxl import load_workbook
path="../test_data/Keen_Testdata.xlsx"
readdata_tasks=Excel_data(path, "Tasks")
from utilities.base_test import BaseTest
readdata_plansassociate=Excel_data(path, "Plans_associate")
from utilities import logger_utils
from Pages.Keen_Task import New_Task
from Pages.Keen_create_Capture_full_Member_details import Create_capture_member_details
Number=BaseTest()
sheet_tasks=Number.get_random_number('Tasks')
class Test_New_Task():
    log = logger_utils.get_logger()
    def __init__(self,driver):
        self.driver=driver
        workbox = load_workbook(path)
        sheet_tasks = workbox['Tasks']
        sheet=sheet_tasks
        self.sheet=sheet
    @allure.description("Create and Associate a Task to Member")
    def test_task(self,row_number):
        keen = New_Task(self.driver)
        action = ActionChains(self.driver)
        keen_cerate = Create_capture_member_details(self.driver)
        action_utils = ActionUtils(self.driver)
        print("Row number:", row_number)
        data_row = self.sheet[row_number]
        column_field_mapping = {
            'Subject': 'Subject',
            'Status': 'Status', 'Dispositions': 'Dispositions', 'Start Date': 'Start Date',
            'Start Time': 'Start Time', 'End Date': 'End Date', 'End Time': 'End Time',
            'Satisfaction Score': 'Satisfaction Score', 'Task Type': 'Task Type', 'Task Subtype': 'Task Subtype',
        }
        field_locators = {
            'Subject': (New_Task.task_subject),
            'Status': (New_Task.task_status),
            'Dispositions': (New_Task.dispositions),
            'Start Date': (New_Task.task_start_date),
            'Start Time': (New_Task.task_start_time),
            'End Date': (New_Task.task_end_date),
            'End Time': (New_Task.task_end_time),
            'Satisfaction Score': (New_Task.select_satisfactionscore),
            'Task Type': (New_Task.task_type),
            'Task Subtype': (New_Task.task_subtype),
        }
        task_name=None
        for cell in data_row:
            column_name = self.sheet.cell(row=2, column=cell.column).value
            if column_name in column_field_mapping:
                field_name = column_field_mapping[column_name]
                column_value = cell.value
                self.log.info("***** Column name and field name is matched *****")
                locator = field_locators[field_name]
                element = self.driver.find_element(*locator)
                action_utils.wait_for_element((locator))
                self.driver.execute_script("arguments[0].scrollIntoView();", element)
                if element.tag_name == 'div':
                    if field_name == 'Dispositions':
                        self.log.info("***** Dispositions field is displayed *****")
                        action_utils.wait_for_elements((New_Task.drop_down_options))
                        options = keen.pick_options_list()
                        for dispositions in options:
                            if dispositions.get_attribute('title') == column_value:
                                self.driver.execute_script("arguments[0].click();", dispositions)
                                self.log.info("***** Disposition option is selected *****")
                                keen.Click_move_to_choosen()
                                self.log.info("***** Option is moved to choosen *****")
                                break
                elif element.tag_name == 'a':
                    if field_name == 'Status':
                        self.driver.execute_script("arguments[0].click();", element)
                        self.log.info("***** Status field is clicked *****")
                        action_utils.wait_for_elements((New_Task.subtask_options))
                        list_status = keen.select_subtask_options()
                        for status in list_status:
                            if status.get_attribute('title') == column_value:
                                print("Status value is matched")
                                self.driver.execute_script("arguments[0].click();", status)
                                self.log.info("***** Status option is clicked *****")
                                break
                    if field_name == 'Task Type':
                        self.driver.execute_script("arguments[0].click();", element)
                        self.log.info("***** Task type field is clicked ****")
                        action_utils.wait_for_elements((New_Task.subtask_options))
                        options = keen.select_subtask_options()
                        for task_type in options:
                            if task_type.get_attribute('title').strip() == column_value.strip():
                                print("Task type value is matched")
                                self.driver.execute_script("arguments[0].click();", task_type)
                                self.log.info("***** Task type option is selected *****")
                                break
                    if field_name == 'Task Subtype':
                        self.driver.execute_script("arguments[0].click();", element)
                        self.log.info("***** Task subtype field is clicked *****")
                        action_utils.wait_for_elements((New_Task.subtask_options))
                        subtask = keen.select_subtask_options()
                        for subtasks in subtask:
                            print("subtasks options:",subtasks.get_attribute('title'))
                            if subtasks.get_attribute('title').strip() == column_value.strip():
                                print("subtasks value is matched")
                                self.driver.execute_script("arguments[0].click();", subtasks)
                                self.log.info("***** Task subtype option is selected *****")
                                break
                    if field_name == 'Satisfaction Score':
                        self.driver.execute_script("arguments[0].click();", element)
                        self.log.info("***** Satisfaction Score field is clicked *****")
                        action_utils.wait_for_elements((New_Task.subtask_options))
                        score = keen.select_subtask_options()
                        for satisfaction_score in score:
                            print("satisfaction score:",satisfaction_score.get_attribute('title'))
                            print("column value:",column_value)
                            if satisfaction_score.get_attribute('title').strip() == str(column_value):
                                print("score value is matched")
                                self.driver.execute_script("arguments[0].click();", satisfaction_score)
                                self.log.info("***** Satisfaction Score field is clicked *****")
                                break
                else:
                    element.clear()
                    element.send_keys(column_value)
        # keen.save()
        return task_name
    def validation_task(self,row_number,task_name):
        try:
            keen = New_Task(self.driver)
            action = ActionChains(self.driver)
            keen_cerate = Create_capture_member_details(self.driver)
            action_utils = ActionUtils(self.driver)
            Subject = readdata_tasks.readdata(row_number, 1)
            Status = readdata_tasks.readdata(row_number, 2)
            Dispositions = readdata_tasks.readdata(row_number, 3)
            start_date=readdata_tasks.readdata(row_number,4)
            Start_Time = readdata_tasks.readdata(row_number, 5)
            End_Time = readdata_tasks.readdata(row_number, 7)
            end_date = readdata_tasks.readdata(row_number, 6)
            Score = readdata_tasks.readdata(row_number, 8)
            Task_Type = readdata_tasks.readdata(row_number, 9)
            Task_Subtype = readdata_tasks.readdata(row_number, 10)
            self.driver.execute_script("scroll(0, 300);")
            action_utils.wait_for_element((New_Task.task_validation))
            validation=keen.new_task_validation()
            for created_record in validation:
                if created_record.text == task_name:
                    action.move_to_element(created_record).click().perform()
                    break
            associated_task=keen.select_task()
            for task in associated_task:
                if task.text == Subject:
                    action.move_to_element(task).click().perform()
                    break
            self.driver.execute_script("scroll(0, 300);")
            validation_fields=keen.validation_fields()
            for i in validation_fields:
                print(i.text)
                if i.text == "Normal":
                    print("Test Pass")
                else:
                    print("Test fail")
            try:
                verify_task_type=keen.validation_tasktype()
                if verify_task_type.text == Task_Type:
                    assert True
                else:
                    assert False
            except Exception as ex:
                print(ex)
            try:
                verify_startdate_time=keen.validation_start_data_time()
                if verify_startdate_time.text == start_date+' '+Start_Time:
                    assert True
                else:
                    assert False
            except Exception as ex:
                print(ex)
            try:
                verify_enddate_time=keen.validation_end_data_time()
                if verify_enddate_time.text == End_Time+' '+End_Time:
                    assert True
                else:
                    assert False
            except Exception as ex:
                print(ex)
            try:
                verify_due_date=keen.validation_Due_date()
                if verify_due_date.text == end_date:
                    assert True
                else:
                   assert False
            except Exception as ex:
                print(ex)
        except Exception as ex:
            print(ex)

    def test_verify_task(self):
        keen = New_Task(self.driver)
        action = ActionChains(self.driver)
        time.sleep(10)
        # self.keen.click_more_module().click()
        # time.sleep(10)
        keen.click_Task_module()
        time.sleep(10)
        list_tasks=keen.click_task()
        for i in list_tasks:
            print(i.text)
            if i.text == "Call":
                i.click()
                break
    def test_create_new_task(self):
        try:
            keen = New_Task(self.driver)
            action = ActionChains(self.driver)
            keen.click_Task_module()
            time.sleep(5)
            keen.click_on_showmore_button().click()
            keen.click_new_task().click()
            time.sleep(10)
            keen.subject("Send Quote")
            keen.click_task_type().click()
            select_task_type = keen.select_list_options()
            for i in select_task_type:
                if i.text == "Meet in person":
                    i.click()
                    break
            keen.start_date("12/16/2022")
            time.sleep(5)
            keen.End_date("12/17/2022")
            time.sleep(10)
            options=keen.select_option()
            for i in options:
                if i.text == "Appointment no show":
                    i.click()
                    break
            keen.Click_move_to_choosen()
            keen.start_time().clear()
            keen.start_time().send_keys("10:30 AM")
            keen.End_time().clear()
            keen.End_time().send_keys("5:30 PM")
            keen.new_task_subtype().click()
            list_subtype=keen.select_list_options()
            for i in list_subtype:
                if i.text == "Email":
                    i.click()
                    break
            keen.new_tast_status()
            select_status = keen.select_list_options()
            for i in select_status:
                if i.text == "Open":
                    i.click()
                    break
            keen.click_save()
        except Exception as ex:
            print(ex)

    @allure.description("Verfiy Auto generated Tasks")
    def test_autocreated_tasks(self,row_plan,row_plans_Associate):
        keen= New_Task(self.driver)
        action=ActionChains(self.driver)
        log=logger_utils.get_logger()
        try:
            keen_test_create=Test_create_capture_full_member_details(self.driver)
            carrier_member_ID = readdata_plansassociate.readdata(row_plans_Associate, 1)
            HRA_Completed_Date = readdata_plansassociate.readdata(row_plans_Associate, 2)
            # Effective_date = readdata_plansassociate.readdata(row_plans_Associate, 3)
            # Enrollment_confirmation_number = readdata_plansassociate.readdata(row_plans_Associate, 4)
            # Plan_End_date = readdata_plansassociate.readdata(row_plans_Associate, 5)
            # Enrollnment_type = readdata_plansassociate.readdata(row_plans_Associate, 6)
            App_subbmission_date =readdata_plansassociate.readdata(row_plans_Associate, 7)
            # App_approval_date = readdata_plansassociate.readdata(row_plans_Associate, 8)
            keen_test_create.test_select_member_to_associate_Cpature_Full_Member_details()
            keen_test_create.test_create_plans(row_plan,row_plans_Associate)
            keen.close().click()
            self.driver.refresh()
            self.driver.execute_script("scroll(0,400);")
            try:
                auto_task = keen.verify_auto_task_generated()
                if auto_task.text == "20-day follow-up":
                    assert True
                else:
                    assert False
            except Exception as ex:
                print(ex)
            ## 20-day follow-up
            tab_options = keen.tasks_tab()
            for i in tab_options:
                print(i.text)
                if i.text == "Tasks":
                    keen.click_Task_module()
                else:
                    keen.click_more().click()
                    break
                time.sleep(5)
                keen.Click_Tab_task()
                break
            time.sleep(10)
            keen.click_Task_module()
            log.info("***** Validating the Auot Generated Tasks fields *****")
            try:
                start_date=keen.auto_task_start_date()
                print(start_date.text)
                sample_list = []
                for i in start_date.text:
                    sample_list.append(i)
                list = "".join(sample_list)
                valid = list.split(" ")
                valid_date = "".join(valid[0])
                date = valid_date.replace(",", "")
                print(date)
                if date == App_subbmission_date:
                    assert True
                else:
                    assert False
            except Exception as ex:
                print(ex)
            try:
                end_date=keen.auto_task_end_date()
                print(end_date.text)
                sample_list = []
                for i in end_date.text:
                    sample_list.append(i)
                list = "".join(sample_list)
                valid = list.split(" ")
                valid_date = "".join(valid[0])
                date = valid_date.replace(",", "")
                print(date)
            except Exception as ex:
                print(ex)
            try:
                auto_comments=keen.Auto_Commnets()
                if auto_comments.text == "Auto-created task based on app submission date":
                    assert True
                else:
                    assert False
            except Exception as ex:
                print(ex)
            try:
                verify=keen.Auto_priority()
                if verify.text == "Normal":
                    assert True
                else:
                    assert False
            except Exception as ex:
                print(ex)
            try:
                verify_status=keen.Auto_open()
                if verify_status.text == "Open":
                    assert True
                else:
                    assert False
            except Exception as ex:
                print(ex)
            try:
                verify_task_type=keen.Auto_task_subtype()
                if verify_task_type.text == "Call":
                    assert True
                else:
                    assert False
            except Exception as ex:
                print(ex)
        except Exception as ex:
            print(ex)

